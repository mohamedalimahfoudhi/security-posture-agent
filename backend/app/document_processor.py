import os
import PyPDF2
from docx import Document
import openpyxl
from typing import Optional
import time
import logging

logger = logging.getLogger(__name__)

class DocumentProcessor:
    """Process various document types and extract text content"""
    
    @staticmethod
    def extract_text_from_pdf(file_path: str) -> str:
        """Extract text from PDF file"""
        logger.info(f"[PDF EXTRACTION] Starting PDF extraction from: {file_path}")
        start_time = time.time()
        text = ""
        try:
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                num_pages = len(pdf_reader.pages)
                logger.info(f"[PDF EXTRACTION] File has {num_pages} pages")
                
                for page_num, page in enumerate(pdf_reader.pages):
                    page_text = page.extract_text()
                    text += page_text + "\n"
                    if page_num % 10 == 0:
                        logger.info(f"[PDF EXTRACTION] Processed {page_num + 1}/{num_pages} pages")
                
                extraction_time = time.time() - start_time
                logger.info(f"[PDF EXTRACTION] Completed in {extraction_time:.2f}s, extracted {len(text)} characters")
        except Exception as e:
            logger.error(f"[PDF EXTRACTION] Failed: {str(e)}")
            raise Exception(f"PDF extraction failed: {str(e)}")
        return text
    
    @staticmethod
    def extract_text_from_docx(file_path: str) -> str:
        """Extract text from DOCX file"""
        logger.info(f"[DOCX EXTRACTION] Starting DOCX extraction from: {file_path}")
        start_time = time.time()
        text = ""
        try:
            doc = Document(file_path)
            num_paragraphs = len(doc.paragraphs)
            logger.info(f"[DOCX EXTRACTION] File has {num_paragraphs} paragraphs")
            
            for para_num, paragraph in enumerate(doc.paragraphs):
                text += paragraph.text + "\n"
                if para_num % 50 == 0:
                    logger.info(f"[DOCX EXTRACTION] Processed {para_num + 1}/{num_paragraphs} paragraphs")
            
            extraction_time = time.time() - start_time
            logger.info(f"[DOCX EXTRACTION] Completed in {extraction_time:.2f}s, extracted {len(text)} characters")
            logger.info(f"[DOCX EXTRACTION] First 500 chars: {text[:500]}")
        except Exception as e:
            logger.error(f"[DOCX EXTRACTION] Failed: {str(e)}")
            raise Exception(f"DOCX extraction failed: {str(e)}")
        return text
    
    @staticmethod
    def extract_text_from_xlsx(file_path: str) -> str:
        """Extract text from XLSX file"""
        logger.info(f"[XLSX EXTRACTION] Starting XLSX extraction from: {file_path}")
        start_time = time.time()
        text = ""
        try:
            workbook = openpyxl.load_workbook(file_path, data_only=True)
            sheet_names = workbook.sheetnames
            logger.info(f"[XLSX EXTRACTION] File has {len(sheet_names)} sheets")
            
            for sheet_name in sheet_names:
                sheet = workbook[sheet_name]
                logger.info(f"[XLSX EXTRACTION] Processing sheet: {sheet_name}")
                for row in sheet.iter_rows(values_only=True):
                    row_text = " ".join([str(cell) if cell is not None else "" for cell in row])
                    text += row_text + "\n"
            
            extraction_time = time.time() - start_time
            logger.info(f"[XLSX EXTRACTION] Completed in {extraction_time:.2f}s, extracted {len(text)} characters")
        except Exception as e:
            logger.error(f"[XLSX EXTRACTION] Failed: {str(e)}")
            raise Exception(f"XLSX extraction failed: {str(e)}")
        return text
    
    @staticmethod
    def extract_text_from_txt(file_path: str) -> str:
        """Extract text from TXT file"""
        logger.info(f"[TXT EXTRACTION] Starting TXT extraction from: {file_path}")
        start_time = time.time()
        try:
            with open(file_path, 'r', encoding='utf-8') as file:
                text = file.read()
            extraction_time = time.time() - start_time
            logger.info(f"[TXT EXTRACTION] Completed in {extraction_time:.2f}s, extracted {len(text)} characters")
            return text
        except UnicodeDecodeError:
            try:
                with open(file_path, 'r', encoding='latin-1') as file:
                    text = file.read()
                extraction_time = time.time() - start_time
                logger.info(f"[TXT EXTRACTION] Completed in {extraction_time:.2f}s (latin-1 encoding), extracted {len(text)} characters")
                return text
            except Exception as e:
                logger.error(f"[TXT EXTRACTION] Failed: {str(e)}")
                raise Exception(f"TXT extraction failed: {str(e)}")
        except Exception as e:
            logger.error(f"[TXT EXTRACTION] Failed: {str(e)}")
            raise Exception(f"TXT extraction failed: {str(e)}")
    
    @staticmethod
    def extract_text(file_path: str, file_type: str) -> str:
        """Extract text based on file type"""
        logger.info(f"[DOCUMENT PROCESSOR] Extracting text from {file_type} file: {file_path}")
        file_type = file_type.lower()
        
        if file_type == 'pdf':
            return DocumentProcessor.extract_text_from_pdf(file_path)
        elif file_type == 'docx':
            return DocumentProcessor.extract_text_from_docx(file_path)
        elif file_type == 'xlsx':
            return DocumentProcessor.extract_text_from_xlsx(file_path)
        elif file_type == 'txt':
            return DocumentProcessor.extract_text_from_txt(file_path)
        else:
            logger.error(f"[DOCUMENT PROCESSOR] Unsupported file type: {file_type}")
            raise Exception(f"Unsupported file type: {file_type}")
    
    @staticmethod
    def get_file_type(filename: str) -> str:
        """Determine file type from filename"""
        extension = filename.split('.')[-1].lower()
        return extension
